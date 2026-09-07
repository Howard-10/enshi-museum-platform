"""Import an Excel artifact catalog while preserving source-row provenance."""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.db.session import SessionLocal
from app.repositories.catalog_repository import CatalogRepository

EXPECTED_HEADERS = ("时代", "文物名称", "地点", "材质", "图片")
ANNOTATION_NAME = re.compile(r"^[（(].+[）)]$")
EMPTY_MARKERS = {"", "-", "—", "无", "n/a", "na"}


@dataclass(frozen=True)
class ParsedCatalogRow:
    sheet_name: str
    row_number: int
    record_kind: str
    artifact_name: str | None
    era: str | None
    location: str | None
    material: str | None
    content_fingerprint: str
    raw_data: dict[str, Any]


def raw_cell_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def normalized_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if text.lower() in EMPTY_MARKERS:
        return None
    return text or None


def row_fingerprint(
    *, era: str | None, name: str | None, location: str | None, material: str | None
) -> str:
    payload = json.dumps(
        {"era": era, "name": name, "location": location, "material": material},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def header_positions(values: list[Any]) -> dict[str, int] | None:
    normalized = [normalized_text(value) for value in values]
    positions = {
        header: normalized.index(header) for header in EXPECTED_HEADERS if header in normalized
    }
    return positions if len(positions) == len(EXPECTED_HEADERS) else None


def parse_catalog_worksheet(sheet_name: str, rows: list[tuple[Any, ...]]) -> list[ParsedCatalogRow]:
    """Parse one supported catalog sheet into normalized, source-traceable rows."""

    if not rows:
        return []
    positions = header_positions(list(rows[0]))
    if positions is None:
        return []

    parsed: list[ParsedCatalogRow] = []
    for row_number, values in enumerate(rows[1:], start=2):
        raw_data = {
            header: raw_cell_value(values[index]) if index < len(values) else None
            for header, index in positions.items()
        }
        era = normalized_text(raw_data["时代"])
        name = normalized_text(raw_data["文物名称"])
        location = normalized_text(raw_data["地点"])
        material = normalized_text(raw_data["材质"])
        if not any((era, name, location, material)):
            continue

        is_annotation = name is not None and bool(ANNOTATION_NAME.fullmatch(name))
        artifact_name = None if is_annotation else name
        record_kind = "annotation" if is_annotation else "artifact"
        parsed.append(
            ParsedCatalogRow(
                sheet_name=sheet_name,
                row_number=row_number,
                record_kind=record_kind,
                artifact_name=artifact_name,
                era=era,
                location=location,
                material=material,
                content_fingerprint=row_fingerprint(
                    era=era,
                    name=artifact_name,
                    location=location,
                    material=material,
                ),
                raw_data=raw_data,
            )
        )
    return parsed


def parse_catalog_workbook(path: Path) -> list[ParsedCatalogRow]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        records: list[ParsedCatalogRow] = []
        for sheet in workbook.worksheets:
            records.extend(
                parse_catalog_worksheet(sheet.title, list(sheet.iter_rows(values_only=True)))
            )
        return records
    finally:
        workbook.close()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def source_uri(path: Path, knowledge_root: Path) -> str:
    try:
        relative_path = path.resolve().relative_to(knowledge_root.resolve()).as_posix()
    except ValueError:
        relative_path = path.name
    return f"imports/xlsx/{relative_path}"


async def import_catalog(
    *,
    source_path: Path,
    knowledge_root: Path,
    dry_run: bool,
) -> dict[str, Any]:
    rows = parse_catalog_workbook(source_path)
    source_sha256 = sha256_file(source_path)
    report: dict[str, Any] = {
        "generated_at": datetime.now(UTC).isoformat(),
        "source_path": str(source_path),
        "source_uri": source_uri(source_path, knowledge_root),
        "source_sha256": source_sha256,
        "dry_run": dry_run,
        "rows_detected": len(rows),
        "imported": 0,
        "existing": 0,
        "artifact_created": 0,
        "artifact_enriched": 0,
        "record_kinds": Counter(row.record_kind for row in rows),
        "exact_duplicate_candidate_groups": 0,
        "exact_duplicate_candidate_rows": 0,
    }
    duplicate_counts = Counter(
        row.content_fingerprint for row in rows if row.record_kind == "artifact"
    )
    duplicate_groups = [count for count in duplicate_counts.values() if count > 1]
    report["exact_duplicate_candidate_groups"] = len(duplicate_groups)
    report["exact_duplicate_candidate_rows"] = sum(duplicate_groups)

    async with SessionLocal() as session:
        repository = CatalogRepository(session)
        for row in rows:
            existing = await repository.get_by_source_row(
                source_sha256=source_sha256,
                sheet_name=row.sheet_name,
                row_number=row.row_number,
            )
            if existing is not None:
                report["existing"] += 1
                continue
            if dry_run:
                report["imported"] += 1
                continue

            artifact, created, enriched = await repository.get_or_create_artifact(
                name=row.artifact_name,
                era=row.era,
            )
            repository.add_record(
                artifact=artifact,
                source_sha256=source_sha256,
                source_uri=report["source_uri"],
                sheet_name=row.sheet_name,
                row_number=row.row_number,
                record_kind=row.record_kind,
                era=row.era,
                location=row.location,
                material=row.material,
                content_fingerprint=row.content_fingerprint,
                raw_data=row.raw_data,
            )
            report["imported"] += 1
            report["artifact_created"] += int(created)
            report["artifact_enriched"] += int(enriched)
        if not dry_run:
            await session.commit()

    report["record_kinds"] = dict(report["record_kinds"])
    return report


async def main() -> None:
    parser = argparse.ArgumentParser(description="Import a structured artifact catalog workbook.")
    parser.add_argument("source", type=Path, help="One Excel workbook with catalog columns.")
    parser.add_argument(
        "--knowledge-root",
        type=Path,
        default=None,
        help="Root directory used to store a portable source reference.",
    )
    parser.add_argument(
        "--dry-run", action="store_true", help="Inspect without writing PostgreSQL."
    )
    parser.add_argument(
        "--report", type=Path, default=None, help="Optional JSON audit report path."
    )
    args = parser.parse_args()

    if not args.source.is_file() or args.source.suffix.lower() != ".xlsx":
        raise SystemExit("source must be an existing .xlsx file")
    report = await import_catalog(
        source_path=args.source,
        knowledge_root=args.knowledge_root or args.source.parent,
        dry_run=args.dry_run,
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    asyncio.run(main())

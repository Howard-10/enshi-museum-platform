"""Safely import human decisions from the evidence-review workbook."""

from __future__ import annotations

import argparse
import asyncio
import json
import uuid
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.models.core import ArtifactDocumentLink, DocumentEvidenceReview
from app.db.session import SessionLocal

ALLOWED_STATUSES = {"approved", "rejected", "needs_review"}


def decisions_from_sheet(sheet, *, id_header: str) -> list[tuple[str, str, str]]:
    headers = [cell.value for cell in sheet[4]]
    positions = {str(value).strip(): index for index, value in enumerate(headers) if value}
    required = {id_header, "审核结论", "审核备注"}
    missing = required - positions.keys()
    if missing:
        raise ValueError(f"工作表 {sheet.title} 缺少列：{', '.join(sorted(missing))}")
    decisions: list[tuple[str, str, str]] = []
    for row in sheet.iter_rows(min_row=5, values_only=True):
        record_id = row[positions[id_header]]
        decision = row[positions["审核结论"]]
        note = row[positions["审核备注"]]
        if not record_id or decision in (None, ""):
            continue
        normalized = str(decision).strip().lower()
        if normalized not in ALLOWED_STATUSES:
            raise ValueError(f"{sheet.title} 中存在无效审核结论：{decision}")
        decisions.append((str(record_id).strip(), normalized, str(note or "").strip()))
    return decisions


def load_decisions(path: Path) -> tuple[list[tuple[str, str, str]], list[tuple[str, str, str]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        links = decisions_from_sheet(workbook["关联审核"], id_header="关联编号")
        documents = decisions_from_sheet(workbook["文档审核"], id_header="文档编号")
        return links, documents
    finally:
        workbook.close()


async def apply_decisions(
    links: list[tuple[str, str, str]], documents: list[tuple[str, str, str]], *, apply: bool
) -> dict[str, object]:
    async with SessionLocal() as session:
        missing_links: list[str] = []
        missing_documents: list[str] = []
        for record_id, status, note in links:
            item = await session.scalar(
                select(ArtifactDocumentLink).where(ArtifactDocumentLink.id == uuid.UUID(record_id))
            )
            if item is None:
                missing_links.append(record_id)
                continue
            if apply:
                item.review_status = status
                item.review_note = note or None
        for record_id, status, note in documents:
            item = await session.scalar(
                select(DocumentEvidenceReview).where(
                    DocumentEvidenceReview.document_id == uuid.UUID(record_id)
                )
            )
            if item is None:
                missing_documents.append(record_id)
                continue
            if apply:
                item.review_status = status
                item.review_note = note or None
        if apply:
            await session.commit()

    return {
        "mode": "applied" if apply else "dry_run",
        "link_decisions_found": len(links),
        "document_decisions_found": len(documents),
        "link_status_counts": dict(Counter(status for _, status, _ in links)),
        "document_status_counts": dict(Counter(status for _, status, _ in documents)),
        "missing_link_ids": missing_links,
        "missing_document_ids": missing_documents,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Import reviewed evidence workbook decisions.")
    parser.add_argument("workbook", type=Path, help="Edited 文档文物关联审核表.xlsx")
    parser.add_argument("--yes", action="store_true", help="Actually write the reviewed statuses")
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"审核表不存在：{args.workbook}")
    links, documents = load_decisions(args.workbook)
    print(json.dumps(asyncio.run(apply_decisions(links, documents, apply=args.yes)), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

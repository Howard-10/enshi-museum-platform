"""Read Excel DISPIMG cell images without modifying the source workbook."""

from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

DISPIMG_PATTERN = re.compile(r'DISPIMG\(\s*"([^"]+)"', re.IGNORECASE)
IMAGE_MIME_TYPES = {
    ".jpeg": "image/jpeg",
    ".jpg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}


@dataclass(frozen=True)
class EmbeddedImage:
    workbook_relative_path: str
    sheet_name: str
    cell_coordinate: str
    artifact_name: str
    image_part: str
    mime_type: str
    data: bytes

    @property
    def source_reference(self) -> str:
        return f"{self.workbook_relative_path}#{self.sheet_name}!{self.cell_coordinate}"

    @property
    def sha256(self) -> str:
        return hashlib.sha256(self.data).hexdigest()


@dataclass(frozen=True)
class WorkbookImageIssue:
    workbook_relative_path: str
    issue_type: str
    detail: str


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def image_id_to_part(archive: zipfile.ZipFile) -> dict[str, str]:
    """Map Excel DISPIMG IDs to their corresponding ``xl/media`` part."""

    names = set(archive.namelist())
    if "xl/cellimages.xml" not in names or "xl/_rels/cellimages.xml.rels" not in names:
        return {}

    images_root = ET.fromstring(archive.read("xl/cellimages.xml"))
    relationships_root = ET.fromstring(archive.read("xl/_rels/cellimages.xml.rels"))
    relationship_targets = {
        relation.attrib["Id"]: str(PurePosixPath("xl") / relation.attrib["Target"])
        for relation in relationships_root
    }
    mapped: dict[str, str] = {}
    for picture in (node for node in images_root.iter() if local_name(node.tag) == "pic"):
        image_id = None
        relationship_id = None
        for node in picture.iter():
            if local_name(node.tag) == "cNvPr":
                image_id = node.attrib.get("name")
            elif local_name(node.tag) == "blip":
                relationship_id = next(
                    (value for key, value in node.attrib.items() if key.endswith("}embed")),
                    None,
                )
        if image_id and relationship_id and relationship_id in relationship_targets:
            mapped[image_id] = relationship_targets[relationship_id]
    return mapped


def row_artifact_name(values: list[object]) -> str | None:
    """Return the readable row label next to DISPIMG formula cells."""

    for value in values:
        if not isinstance(value, str):
            continue
        cleaned = value.strip()
        if cleaned and "DISPIMG(" not in cleaned.upper():
            return cleaned
    return None


def read_embedded_images(
    workbook_path: Path,
    knowledge_root: Path,
) -> tuple[list[EmbeddedImage], list[WorkbookImageIssue]]:
    """Extract only cell-mapped images and report any unsafe-to-map leftovers."""

    workbook_relative_path = workbook_path.relative_to(knowledge_root).as_posix()
    with zipfile.ZipFile(workbook_path) as archive:
        id_to_part = image_id_to_part(archive)
        media_parts = {
            name
            for name in archive.namelist()
            if name.startswith("xl/media/") and not name.endswith("/")
        }
        if not id_to_part:
            return [], [
                WorkbookImageIssue(
                    workbook_relative_path=workbook_relative_path,
                    issue_type="no_cell_image_mapping",
                    detail=f"{len(media_parts)} embedded image part(s) cannot be linked to a row safely.",
                )
            ]

        workbook = load_workbook(workbook_path, read_only=True, data_only=False)
        requested: list[tuple[str, str, str, str]] = []
        unmapped_formula_ids: set[str] = set()
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                values = [cell.value for cell in row]
                artifact_name = row_artifact_name(values)
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    for image_id in DISPIMG_PATTERN.findall(cell.value):
                        image_part = id_to_part.get(image_id)
                        if artifact_name and image_part in media_parts:
                            requested.append(
                                (sheet.title, cell.coordinate, artifact_name, image_part)
                            )
                        else:
                            unmapped_formula_ids.add(image_id)

        images: list[EmbeddedImage] = []
        referenced_parts: set[str] = set()
        for sheet_name, cell_coordinate, artifact_name, image_part in requested:
            suffix = Path(image_part).suffix.lower()
            mime_type = IMAGE_MIME_TYPES.get(suffix)
            if mime_type is None:
                continue
            referenced_parts.add(image_part)
            images.append(
                EmbeddedImage(
                    workbook_relative_path=workbook_relative_path,
                    sheet_name=sheet_name,
                    cell_coordinate=cell_coordinate,
                    artifact_name=artifact_name,
                    image_part=image_part,
                    mime_type=mime_type,
                    data=archive.read(image_part),
                )
            )

    issues: list[WorkbookImageIssue] = []
    for image_id in sorted(unmapped_formula_ids):
        issues.append(
            WorkbookImageIssue(
                workbook_relative_path=workbook_relative_path,
                issue_type="unmapped_dispimg_id",
                detail=image_id,
            )
        )
    for image_part in sorted(media_parts - referenced_parts):
        issues.append(
            WorkbookImageIssue(
                workbook_relative_path=workbook_relative_path,
                issue_type="unreferenced_embedded_image",
                detail=image_part,
            )
        )
    return images, issues
